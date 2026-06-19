# -*- coding: utf-8 -*-
from docx import Document
from docx.oxml.ns import qn

DOCX = r"C:\Users\alexdu\Downloads\vFlat\家谱\窑湾_关键边框v2.docx"
doc = Document(DOCX)

t = doc.tables[2]
print("Table[2] -- title column (col 23) borders per row:")
for r in range(len(t.rows)):
    tc = t.cell(r, 23)._tc
    tcPr = tc.find(qn("w:tcPr"))
    out = {"top": "?", "left": "?", "bottom": "?", "right": "?"}
    if tcPr is not None:
        tcBorders = tcPr.find(qn("w:tcBorders"))
        if tcBorders is not None:
            for side in out:
                node = tcBorders.find(qn(f"w:{side}"))
                if node is not None:
                    out[side] = node.get(qn("w:val"))
    print(f"  row[{r}] col[23]: {out}")